"""Parse derivatives market data Excel (先物・オプション取引概況).

対象: {yyyymmdd}_derivatives_market_data_whole_day.xlsx の market_data_OP シート。
日経225オプションのブロック（先頭ブロック、実測では行12-15）:
  C列 = セッション区分（夜間/前場/後場/合計）
  D-I = 取引高: PUT / うちJ-NET / CALL / うちJ-NET / 合計 / うちJ-NET（枚）
  J-O = 売買代金: 同構成（円）
"""
from __future__ import annotations

import io
import logging

import openpyxl

logger = logging.getLogger(__name__)

_SESSION_KEYS = ("夜間", "前場", "後場", "合計")


def parse_op_market_data(content: bytes) -> list[dict]:
    """日経225オプションのセッション別 P/C 取引高・売買代金を返す。

    Returns: [{session, put_volume, call_volume, put_value, call_value,
               *_jnet, total_volume, total_value, ...}] （代金は円）
    先頭の商品ブロック（日経225オプション）のみ。合計行で打ち切る。
    """
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    # summary_data_OP（9列・P/C分解なし）を誤って掴まないよう完全一致を優先
    ws = next((s for s in wb.worksheets if s.title == "market_data_OP"), None)
    if ws is None:
        ws = next((s for s in wb.worksheets
                   if "market_data" in s.title and "OP" in s.title), None)
    if ws is None:
        logger.warning("market_data_OP sheet not found (sheets=%s)",
                       [s.title for s in wb.worksheets])
        wb.close()
        return []

    rows: list[dict] = []
    in_block = False
    for r in range(5, 46):
        label = ws.cell(row=r, column=3).value
        put_val = ws.cell(row=r, column=10).value  # J: PUT売買代金
        is_sess = isinstance(label, str) and any(k in label for k in _SESSION_KEYS)
        if is_sess and isinstance(put_val, (int, float)):
            if not in_block:
                # ブロック先頭で商品名を検証（日経225オプション・ラージのみ対象）
                product = str(ws.cell(row=r, column=2).value or "")
                if "日経225オプション" not in product or "ミニ" in product:
                    logger.warning("unexpected first OP block: %r", product[:30])
                    break
            in_block = True

            def _n(col: int) -> float:
                v = ws.cell(row=r, column=col).value
                return float(v) if isinstance(v, (int, float)) else 0.0

            rows.append({
                "session": label.strip(),
                "put_volume": _n(4), "put_volume_jnet": _n(5),
                "call_volume": _n(6), "call_volume_jnet": _n(7),
                "total_volume": _n(8), "total_volume_jnet": _n(9),
                "put_value": _n(10), "put_value_jnet": _n(11),
                "call_value": _n(12), "call_value_jnet": _n(13),
                "total_value": _n(14), "total_value_jnet": _n(15),
            })
            if "合計" in label:
                break  # 日経225ブロック終端（以降は別商品）
        elif in_block:
            break
    wb.close()
    return rows
