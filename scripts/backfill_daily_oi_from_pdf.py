"""JPX日報PDF（Daily_Report_OSE_{yyyymmdd}.zip）から建玉残高表・取引概況xlsxを復元する。

JPXは建玉残高表（{d}open_interest.xlsx）と取引概況（{d}_derivatives_market_data_
whole_day.xlsx）を最新営業日分しか掲載しない。収集漏れ日（2026-09-30, 10-01: 添付ID
変更）はExcelが永久欠損になるため、過去分も残る日報PDFから既存パーサ
（data/parser_daily_oi.py, data/parser_market_data.py）がそのまま読めるxlsxを再構成する。

復元範囲（2026-09-29 実ファイルとの突合で検証）:
  sheet0  先物6商品（日経225先物/TOPIX先物/日経225mini/日経225マイクロ/ミニTOPIX/JPX日経400）
          取引高 = 競争売買 + J-NET、当日建玉残高 = PDF建玉残高、前日建玉残高 = 前営業日ファイル
  別紙1   日経225オプション 直近3限月（取引高 = 競争売買 + J-NET）
  取引概況 market_data_OP の「合計」行のみ（P/C 取引高・売買代金とうちJ-NET）
  セッション別内訳・別紙2（ミニオプション）・上記以外の商品は復元しない。

使い方（日付は昇順に処理。前日建玉は直前に復元した日も参照する）:
  python scripts/backfill_daily_oi_from_pdf.py --selftest
  python scripts/backfill_daily_oi_from_pdf.py 20260930 20261001            # out へ書き出しのみ
  python scripts/backfill_daily_oi_from_pdf.py 20260930 20261001 --upload   # キャッシュ(L1+R2)へ保存
"""
from __future__ import annotations

import argparse
import io
import sys
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import requests

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import config  # noqa: E402
from data import fetcher  # noqa: E402
from data.cache import ensure_cache_dirs, save_to_cache  # noqa: E402
from data.parser_daily_oi import parse_daily_futures_oi_excel, parse_daily_oi_excel  # noqa: E402
from data.parser_daily_report_futures import parse_sif_pdf  # noqa: E402
from data.parser_daily_report_options import parse_siop_pdf  # noqa: E402
from data.parser_market_data import parse_op_market_data  # noqa: E402

ZIP_URL = ("https://www.jpx.co.jp/automation/markets/statistics-derivatives/daily/files/"
           "{yyyymm}/Daily_Report_OSE_{yyyymmdd}.zip")
ZIP_DIR = config.CACHE_DIR / "daily_report"
TAG = "【日報PDFより復元】"
LEFT = [("日経225先物", "NK225F"), ("TOPIX先物", "TOPIXF")]          # parser: 列A 部分一致
RIGHT = [("日経225mini", "NK225MF"), ("日経225マイクロ", "NK225MicroF"),
         ("ミニTOPIX", "MiniTOPIXF"), ("JPX日経400", "JPX400F")]      # parser: 列H 部分一致
HDR = ["限月取引", "取引高", "当日\n建玉残高", "前日比", "前日\n建玉残高"]
FIELDS = ("trading_volume", "current_oi", "net_change", "previous_oi")


def get_zip(d: date) -> bytes | None:
    """日報zip。ローカル保存分を優先、無ければJPXから取得（未公表は None）。"""
    ds = d.strftime("%Y%m%d")
    p = ZIP_DIR / f"Daily_Report_OSE_{ds}.zip"
    if p.exists():
        return p.read_bytes()
    r = requests.get(ZIP_URL.format(yyyymm=ds[:6], yyyymmdd=ds), timeout=120,
                     headers={"User-Agent": "Mozilla/5.0"})
    if r.status_code != 200 or not r.content.startswith(b"PK"):
        return None
    ZIP_DIR.mkdir(parents=True, exist_ok=True)
    p.write_bytes(r.content)
    return r.content


def parse_zip(d: date) -> tuple[list[dict], list[dict]] | None:
    """日報zip → (先物行, 日経225オプション行)。未公表は None。"""
    z = get_zip(d)
    if z is None:
        return None
    ds = d.strftime("%Y%m%d")
    with zipfile.ZipFile(io.BytesIO(z)) as zf:
        d1, fut = parse_sif_pdf(zf.read(f"sif_dyr_{ds}.pdf"))
        d2, opt = parse_siop_pdf(zf.read(f"siop_dyr_{ds}.pdf"))
    if d1 != d or d2 != d:
        raise ValueError(f"PDF日付不一致: {d} sif={d1} siop={d2}")
    return fut, opt


def _first_fut(content: bytes) -> dict:
    """アプリ（load_daily_futures_oi）と同じ先頭一致で (product, cm) → record。"""
    out: dict = {}
    for r in parse_daily_futures_oi_excel(content):
        out.setdefault((r.product, r.contract_month), r)
    return out


def _opt_map(content: bytes) -> dict:
    return {(r.contract_month, r.option_type, r.strike_price): r for r in parse_daily_oi_excel(content)}


def _cm_label(cm: str) -> str:
    return f"{2000 + int(cm[:2])}年{cm[2:]}月限"


def build_oi_xlsx(d: date, fut: list[dict], opt: list[dict], prev: bytes | None) -> bytes:
    """建玉残高表xlsxを復元（sheet0=先物、別紙1=日経225オプション、別紙2=空）。"""
    pf = _first_fut(prev) if prev else {}
    po = _opt_map(prev) if prev else {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "デリバティブ建玉残高状況"
    ws["A1"] = "JPXデリバティブ建玉残高状況（概算）" + TAG
    ws["A2"] = datetime(d.year, d.month, d.day)
    ws["A4"] = "＜OSE 指数先物取引＞"
    for i, h in enumerate(HDR):
        ws.cell(5, 2 + i, h)
        ws.cell(5, 9 + i, h)
    by_prod: dict[str, dict[str, tuple[int, int]]] = {}
    for r in fut:
        if r.get("product"):
            vol = int(r.get("auction_volume") or 0) + int(r.get("jnet_volume") or 0)
            by_prod.setdefault(r["product"], {})[r["contract_month"]] = (vol, int(r["open_interest"] or 0))
    for products, c0 in ((LEFT, 1), (RIGHT, 8)):
        row = 6
        for name, code in products:
            months = sorted(by_prod.get(code, {}))
            if not months:
                continue
            tot = [0, 0, 0, 0]
            for j, cm in enumerate(months):
                vol, oi = by_prod[code][cm]
                p = pf.get((code, cm))
                prev_oi = p.current_oi if p else 0
                if j == 0:
                    ws.cell(row, c0, name)
                ws.cell(row, c0 + 1, _cm_label(cm))
                for k, v in enumerate([vol, oi, oi - prev_oi, prev_oi]):
                    ws.cell(row, c0 + 2 + k, v)
                    tot[k] += v
                row += 1
            ws.cell(row, c0 + 1, "合計")
            for k, v in enumerate(tot):
                ws.cell(row, c0 + 2 + k, v)
            row += 1

    # 別紙1: 直近3限月。取引高は市場合算、建玉は競争売買行（無ければJ-NET行）
    agg: dict[tuple[str, str, int], list] = {}
    for r in opt:
        k = (r["contract_month"], r["option_type"], int(r["strike_price"]))
        a = agg.setdefault(k, [0, None])
        a[0] += int(r["trading_volume"] or 0)
        if r["market"] == "auction" or a[1] is None:
            a[1] = int(r["open_interest"] or 0)
    months = sorted({k[0] for k, a in agg.items() if a[1]})[:3]
    ws1 = wb.create_sheet("別紙1")
    ws1["A1"] = "JPX デリバティブ建玉残高状況（別紙1）" + TAG
    ws1["A2"] = datetime(d.year, d.month, d.day)
    ws1["A5"], ws1["G5"] = "＜日経225オプション（プット）＞", "＜日経225オプション（コール）＞"
    for i, h in enumerate(HDR):
        ws1.cell(6, 1 + i, h)
        ws1.cell(6, 7 + i, h)
    keys = {k for k in agg if k[0] in months} | {k for k in po if k[0] in months}
    for ot, c0, lab in (("PUT", 1, "プット合計"), ("CALL", 7, "コール合計")):
        row, tot = 7, [0, 0, 0, 0]
        for k in sorted(k for k in keys if k[1] == ot):
            vol, oi = agg.get(k, [0, 0])
            oi = oi or 0
            p = po.get(k)
            prev_oi = p.current_oi if p else 0
            if not (vol or oi or prev_oi):
                continue  # 実ファイルは建玉・取引高とも0の銘柄を載せない
            ws1.cell(row, c0, f"NIKKEI 225 {ot[0]}{k[0]}-{k[2]}")
            for j, v in enumerate([vol, oi, oi - prev_oi, prev_oi]):
                ws1.cell(row, c0 + 1 + j, v)
                tot[j] += v
            row += 1
        ws1.cell(row, c0, lab)
        for j, v in enumerate(tot):
            ws1.cell(row, c0 + 1 + j, v)
    ws2 = wb.create_sheet("別紙2")
    ws2["A1"] = "JPX デリバティブ建玉残高状況（別紙2）" + TAG + "※未復元"
    ws2["A2"] = datetime(d.year, d.month, d.day)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_market_xlsx(d: date, opt: list[dict]) -> bytes:
    """取引概況xlsx（market_data_OP の合計行のみ）を復元。列D-O は実ファイルと同じ並び。"""
    s: dict[str, int] = {}
    for r in opt:
        for key, v in ((f"{r['option_type']}_vol", r["trading_volume"]),
                       (f"{r['option_type']}_val", r["trading_value_yen"])):
            s[key] = s.get(key, 0) + int(v or 0)
            if r["market"] == "jnet":
                s[key + "_j"] = s.get(key + "_j", 0) + int(v or 0)
    g = s.get
    vals = [g("PUT_vol", 0), g("PUT_vol_j", 0), g("CALL_vol", 0), g("CALL_vol_j", 0),
            g("PUT_vol", 0) + g("CALL_vol", 0), g("PUT_vol_j", 0) + g("CALL_vol_j", 0),
            g("PUT_val", 0), g("PUT_val_j", 0), g("CALL_val", 0), g("CALL_val_j", 0),
            g("PUT_val", 0) + g("CALL_val", 0), g("PUT_val_j", 0) + g("CALL_val_j", 0)]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "market_data_OP"
    ws["A1"] = f"デリバティブ取引市況 {d:%Y-%m-%d}" + TAG + "（合計行のみ）"
    ws.cell(5, 2, "日経225オプション")
    ws.cell(5, 3, "合計")
    for i, v in enumerate(vals):
        ws.cell(5, 4 + i, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _prev_file(d: date, built: dict[date, bytes]) -> tuple[date, bytes] | None:
    """直前の営業日の建玉残高表（今回復元分 → キャッシュ/JPX の順）。"""
    for i in range(1, 8):
        p = d - timedelta(days=i)
        if p in built:
            return p, built[p]
        if p.weekday() < 5:
            b = fetcher.download_daily_oi_excel(p)
            if b is not None:
                return p, b
    return None


def backfill(dates: list[date], out_dir: Path, upload: bool, force: bool = False) -> dict[date, bytes]:
    ensure_cache_dirs()
    out_dir.mkdir(parents=True, exist_ok=True)
    built: dict[date, bytes] = {}
    for d in sorted(dates):
        ds = d.strftime("%Y%m%d")
        oi_fn = config.DAILY_OI_FILENAME.format(yyyymmdd=ds)
        md_fn = config.MARKET_DATA_FILENAME.format(yyyymmdd=ds, session="whole_day")
        need_oi = force or fetcher._cached_any(oi_fn, config.CACHE_DAILY_OI_DIR, 24 * 3650) is None
        need_md = force or fetcher.cached_market_data_excel(d) is None
        if not (need_oi or need_md):
            print(f"{ds}: 実ファイルあり、スキップ")
            continue
        parsed = parse_zip(d)
        if parsed is None:
            print(f"{ds}: 日報zip未公表")
            continue
        fut, opt = parsed
        if need_oi:
            prev = _prev_file(d, built)
            oi_x = build_oi_xlsx(d, fut, opt, prev[1] if prev else None)
            built[d] = oi_x
            (out_dir / oi_fn).write_bytes(oi_x)
            print(f"{ds}: 建玉復元 prev={prev[0] if prev else None} "
                  f"先物{len(parse_daily_futures_oi_excel(oi_x))}行 OP{len(parse_daily_oi_excel(oi_x))}行")
            if upload:
                save_to_cache(fetcher._canonical_url(oi_fn), config.CACHE_DAILY_OI_DIR, oi_x)
        if need_md:
            md_x = build_market_xlsx(d, opt)
            (out_dir / md_fn).write_bytes(md_x)
            print(f"{ds}: 取引概況(P/C代金)復元")
            if upload:
                save_to_cache(fetcher._canonical_url(md_fn), config.CACHE_MARKET_DATA_DIR, md_x)
        if upload:
            print(f"{ds}: キャッシュ(L1+R2)へ保存")
    return built


def _cmp(a: dict, b: dict) -> tuple[int, list]:
    both = a.keys() & b.keys()
    bad = [(k, [getattr(a[k], f) for f in FIELDS], [getattr(b[k], f) for f in FIELDS])
           for k in both if any(getattr(a[k], f) != getattr(b[k], f) for f in FIELDS)]
    return len(both), bad + [(k, "片側のみ") for k in a.keys() ^ b.keys()]


def selftest(out_dir: Path) -> None:
    """(a) 9/29 を復元し実ファイルと全項目突合 (b) 復元10/1の建玉 vs 実10/2の前日建玉。"""
    f6 = {c for _, c in LEFT + RIGHT}
    d29 = date(2026, 9, 29)
    real29 = fetcher.download_daily_oi_excel(d29)
    fut29, opt29 = parse_zip(d29)
    built29 = build_oi_xlsx(d29, fut29, opt29, _prev_file(d29, {})[1])
    n, bad = _cmp({k: v for k, v in _first_fut(built29).items() if k[0] in f6},
                  {k: v for k, v in _first_fut(real29).items() if k[0] in f6})
    print(f"(a) 9/29 先物 {n}件比較 不一致{len(bad)}", bad[:5])
    n, bad = _cmp(_opt_map(built29), _opt_map(real29))
    print(f"(a) 9/29 OP別紙1 {n}件比較 不一致{len(bad)}", bad[:5])
    t_b = next(r for r in parse_op_market_data(build_market_xlsx(d29, opt29)) if "合計" in r["session"])
    t_r = next(r for r in parse_op_market_data(fetcher.cached_market_data_excel(d29)) if "合計" in r["session"])
    print("(a) 9/29 取引概況 不一致:", {k: (t_b[k], t_r[k]) for k in t_r if k != "session" and t_b.get(k) != t_r[k]})
    built = backfill([date(2026, 9, 30), date(2026, 10, 1)], out_dir, upload=False)
    b01, r02 = built[date(2026, 10, 1)], fetcher.download_daily_oi_excel(date(2026, 10, 2))
    fa, fb = _first_fut(b01), _first_fut(r02)
    ks = [k for k in fa.keys() & fb.keys() if k[0] in f6]
    print(f"(b) 先物 {len(ks)}件 不一致", [(k, fa[k].current_oi, fb[k].previous_oi) for k in ks
                                     if fa[k].current_oi != fb[k].previous_oi][:5])
    oa, ob = _opt_map(b01), _opt_map(r02)
    ks = oa.keys() & ob.keys()
    print(f"(b) OP {len(ks)}件 不一致", [(k, oa[k].current_oi, ob[k].previous_oi) for k in ks
                                    if oa[k].current_oi != ob[k].previous_oi][:5],
          "10/2側のみ(前日建玉>0)", [k for k in ob.keys() - oa.keys() if ob[k].previous_oi][:5])


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("dates", nargs="*", help="yyyymmdd")
    ap.add_argument("--upload", action="store_true")
    ap.add_argument("--force", action="store_true", help="実ファイルがあっても上書き")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--out", default=str(config.CACHE_DIR / "backfill_out"))
    a = ap.parse_args()
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    if a.selftest:
        selftest(Path(a.out))
    else:
        backfill([datetime.strptime(s, "%Y%m%d").date() for s in a.dates], Path(a.out), a.upload, a.force)
